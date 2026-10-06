"""
Build a cover letter .docx from scratch, styled to match a resume's
extracted font/color profile.

Unlike `docx_rewrite.py` (used by `cv-translate-de` and `resume-tailor`),
this does not modify an existing document's paragraph slots — there is no
pre-existing cover letter to preserve the formatting of. It authors a new,
simple business-letter document (plain paragraphs, no bullets, no tables)
using style values pulled from the actual resume via `docx_inspect.py`, so
the two documents look like one coordinated package. See
`references/visual-consistency.md` for why building fresh is safe here.

Plan format (JSON), one object with "style" and "blocks":

{
  "style": {
    "heading_font": "Calibri", "heading_size_pt": 20,
    "heading_color": "1F3864", "heading_bold": true,
    "body_font": "Calibri", "body_size_pt": 11,
    "muted_color": "555555",
    "accent_color": "156082"
  },
  "blocks": [
    {"role": "sender", "lines": ["Firstname Lastname", "City, Country", "email | phone"]},
    {"role": "recipient", "lines": ["Firma GmbH", "Frau Nachname", "Strasse Nr.", "PLZ Ort"]},
    {"role": "date_place", "text": "Wuerzburg, 27. September 2026"},
    {"role": "subject", "text": "Bewerbung als AI Engineer, Kennziffer 12345"},
    {"role": "salutation", "text": "Sehr geehrte Frau Nachname,"},
    {"role": "body", "segments": [{"text": "..."}, {"text": "30 %", "bold": true}, {"text": "..."}]},
    {"role": "body", "segments": [{"text": "..."}]},
    {"role": "closing_phrase", "text": "Mit freundlichen Gruessen"},
    {"role": "signature_name", "text": "Firstname Lastname"},
    {"role": "enclosures", "text": "Anlagen: Lebenslauf"}
  ]
}

Roles: sender, recipient, date_place (right-aligned), subject (bold),
salutation, body (repeatable, supports inline bold segments the same way
`docx_rewrite.py` plans do), closing_phrase, signature_name, enclosures
(rendered in muted_color if given).

Style keys are all optional; anything omitted falls back to a plain
default so the script still produces a readable letter without a full
style profile.

Usage:
    python build_cover_letter_docx.py <plan.json> <output.docx>
"""
import json
import sys

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

DEFAULT_BODY_FONT = "Calibri"
DEFAULT_BODY_SIZE = 11.0


def _set_font(run, font_name=None, size_pt=None, color=None, bold=None):
    if font_name:
        run.font.name = font_name
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.font.bold = bold


def build(plan, output_path):
    style = plan.get("style", {})
    body_font = style.get("body_font", DEFAULT_BODY_FONT)
    body_size = style.get("body_size_pt", DEFAULT_BODY_SIZE)
    heading_font = style.get("heading_font", body_font)
    heading_size = style.get("heading_size_pt", body_size + 6)
    heading_color = style.get("heading_color")
    heading_bold = style.get("heading_bold", True)
    muted_color = style.get("muted_color")

    d = docx.Document()
    normal = d.styles["Normal"]
    normal.font.name = body_font
    normal.font.size = Pt(body_size)
    # python-docx's built-in default template adds 10pt space-after and
    # 1.15 line spacing to every paragraph (see w:pPrDefault in its
    # styles.xml). Spacing between blocks here is controlled explicitly via
    # blank paragraphs, so that default compounds into pages of dead space
    # if left in place.
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    def add_line(text, font_name=body_font, size_pt=body_size, color=None,
                 bold=None, alignment=None):
        p = d.add_paragraph()
        if alignment is not None:
            p.alignment = alignment
        run = p.add_run(text)
        _set_font(run, font_name, size_pt, color, bold)
        return p

    def add_segments(segments, alignment=None):
        p = d.add_paragraph()
        if alignment is not None:
            p.alignment = alignment
        for seg in segments:
            run = p.add_run(seg.get("text", ""))
            _set_font(
                run,
                body_font,
                body_size,
                seg.get("color"),
                seg.get("bold"),
            )
        return p

    def blank():
        d.add_paragraph()

    for block in plan.get("blocks", []):
        role = block["role"]
        if role == "sender":
            lines = block["lines"]
            if lines:
                add_line(lines[0], heading_font, heading_size, heading_color, heading_bold)
                for line in lines[1:]:
                    add_line(line)
            blank()
        elif role == "recipient":
            for line in block["lines"]:
                add_line(line)
            blank()
        elif role == "date_place":
            add_line(block["text"], alignment=WD_ALIGN_PARAGRAPH.RIGHT)
            blank()
        elif role == "subject":
            add_line(block["text"], bold=True)
            blank()
        elif role == "salutation":
            add_line(block["text"])
            blank()
        elif role == "body":
            add_segments(block["segments"])
            blank()
        elif role == "closing_phrase":
            add_line(block["text"])
            blank()
        elif role == "signature_name":
            add_line(block["text"])
            blank()
        elif role == "enclosures":
            add_line(block["text"], size_pt=body_size - 1 if body_size > 8 else body_size,
                      color=muted_color)
        else:
            raise ValueError(f"Unknown block role: {role!r}")

    d.save(output_path)
    print(f"Wrote {output_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python build_cover_letter_docx.py <plan.json> <output.docx>", file=sys.stderr)
        sys.exit(2)
    with open(sys.argv[1], encoding="utf-8") as f:
        plan_data = json.load(f)
    build(plan_data, sys.argv[2])
