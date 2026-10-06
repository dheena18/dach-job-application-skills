"""
Apply a rewrite plan to a .docx, replacing the runs of specific paragraphs
while keeping every other paragraph (and the document's styles, theme,
margins, section layout) completely untouched.

A paragraph is rewritten by deleting all of its <w:r> runs and adding new
ones from the plan's "segments" list. This is deliberate: it avoids the
classic "find matching substring across fragmented runs" problem entirely,
because we are not searching for old text inside old runs, we are replacing
the paragraph's whole run sequence with a new one that we authored. Anything
about the paragraph itself (alignment, tab stops, bullet numbering, spacing)
lives in <w:pPr>, which this script never touches.

Plan format (JSON), a list of objects:
[
  {
    "index": 11,
    "segments": [
      {"text": "Trainierte ein UNet-Diffusionsmodell", "bold": true},
      {"text": " unter Verwendung von ", "bold": false},
      {"text": "Classifier-Free Guidance", "bold": true},
      {"text": " ..."}
    ]
  },
  ...
]

Each segment may set: text (required), bold, italic, color (hex string, no
"#"), size_pt (float). Omitted keys leave that formatting attribute unset on
the new run (i.e. inherited from the paragraph/document style, matching how
the original unformatted runs behaved).

Usage:
    python docx_rewrite.py <input.docx> <plan.json> <output.docx>
"""
import json
import sys

import docx
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def _clear_runs(paragraph):
    p = paragraph._p
    for child in list(p):
        if child.tag == qn("w:r"):
            p.remove(child)
        elif child.tag == qn("w:hyperlink"):
            raise RuntimeError(
                f"Paragraph {paragraph.text!r} contains a <w:hyperlink> run; "
                "this script does not rewrite hyperlink runs. Handle it manually."
            )


def _apply_segment(paragraph, seg):
    run = paragraph.add_run(seg.get("text", ""))
    if "bold" in seg and seg["bold"] is not None:
        run.font.bold = seg["bold"]
    if "italic" in seg and seg["italic"] is not None:
        run.font.italic = seg["italic"]
    if seg.get("color"):
        run.font.color.rgb = RGBColor.from_string(seg["color"])
    if seg.get("size_pt") is not None:
        run.font.size = Pt(seg["size_pt"])


def rewrite(input_path, plan_path, output_path):
    d = docx.Document(input_path)
    with open(plan_path, encoding="utf-8") as f:
        plan = json.load(f)

    paragraphs = d.paragraphs
    touched = set()
    for entry in plan:
        idx = entry["index"]
        if idx in touched:
            raise ValueError(f"Paragraph {idx} appears twice in the plan.")
        touched.add(idx)
        if idx >= len(paragraphs):
            raise IndexError(f"Paragraph index {idx} out of range (doc has {len(paragraphs)}).")
        para = paragraphs[idx]
        _clear_runs(para)
        for seg in entry["segments"]:
            _apply_segment(para, seg)

    d.save(output_path)
    print(f"Rewrote {len(touched)} paragraph(s). Saved to {output_path}")


if __name__ == "__main__":
    if len(sys.argv) != 4:
        print("Usage: python docx_rewrite.py <input.docx> <plan.json> <output.docx>", file=sys.stderr)
        sys.exit(2)
    rewrite(sys.argv[1], sys.argv[2], sys.argv[3])
